"""General Reports viewers can open the users the reports page lists."""

import io
import json

from dlux.tests.harness import setup_test_environment

setup_test_environment()

from django.apps import apps
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import Client, TestCase
from django.urls import reverse

from dlux.utils import USER_REPORT_ACTIVITY, USER_REPORT_FULL, user_report_access

User = get_user_model()


class ReportViewerUserReportTests(TestCase):
    def setUp(self):
        SystemSettings = apps.get_model('dlux', 'SystemSettings')
        settings_obj = SystemSettings.load()
        settings_obj.is_configured = True
        settings_obj.save(update_fields=['is_configured'])
        self.ActivityLog = apps.get_model('dlux', 'ActivityLog')
        self.Scope = apps.get_model('dlux', 'Scope')
        self.viewer = User.objects.create_user('central', password='centralpass1', is_staff=True)
        self.viewer.user_permissions.add(Permission.objects.get(codename='view_reports', content_type__app_label='dlux'))
        self.worker = User.objects.create_user('worker', email='worker@example.com', password='x')
        for action in ('CREATE', 'UPDATE', 'UPDATE', 'DELETE'):
            self.ActivityLog.objects.create(
                created_by=self.worker, action=action, model_name='Project Entry',
                model_key='tests.entry', ip_address='203.0.113.77', user_agent='SecretAgent/1.0',
            )
        self.client = Client()
        self.client.login(username='central', password='centralpass1')

    def _modal(self, user, **params):
        response = self.client.get(
            reverse('user_report_modal', args=[user.pk]), {'window': 'all', **params},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        return response

    def _enable_scopes(self):
        ScopeSettings = apps.get_model('dlux', 'ScopeSettings')
        settings_obj = ScopeSettings.load()
        settings_obj.is_enabled = True
        settings_obj.save()

    def test_report_viewer_gets_the_activity_breakdown_not_403(self):
        self.assertEqual(user_report_access(self.viewer, self.worker), USER_REPORT_ACTIVITY)
        response = self._modal(self.worker)
        self.assertEqual(response.status_code, 200)
        html = json.loads(response.content)['html']
        self.assertIn('Project Entry', html)
        self.assertIn('data-dlux-user-report-activity-item', html)
        # Nothing beyond activity leaks into the reports view.
        for secret in ('203.0.113.77', 'SecretAgent', 'worker@example.com'):
            self.assertNotIn(secret, html)
        self.assertNotIn('dlux-user-report-devices', html)

    def test_activity_export_has_no_network_columns(self):
        from openpyxl import load_workbook

        response = self.client.get(reverse('user_report_xlsx', args=[self.worker.pk]), {'window': 'all'})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(len(workbook.sheetnames), 3)
        values = {
            str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
            if cell.value is not None
        }
        self.assertFalse({'203.0.113.77', 'SecretAgent/1.0', 'worker@example.com'} & values)

    def test_scoped_users_stay_out_of_a_central_viewers_reach(self):
        self._enable_scopes()
        branch = self.Scope.objects.create(name='Branch')
        scoped = User.objects.create_user('scoped', password='x')
        scoped.profile.scope = branch
        scoped.profile.save()
        self.assertIsNone(user_report_access(self.viewer, scoped))
        self.assertEqual(self._modal(scoped).status_code, 403)
        # The worker is unscoped, so central staff still sees them.
        self.assertEqual(self._modal(self.worker).status_code, 200)

    def test_scope_limits_the_actions_counted_like_the_reports_page(self):
        self._enable_scopes()
        branch = self.Scope.objects.create(name='Branch')
        self.ActivityLog.objects.create(created_by=self.worker, action='CREATE', model_name='Branch Entry', scope=branch)
        html = json.loads(self._modal(self.worker).content)['html']
        self.assertIn('Project Entry', html)
        self.assertNotIn('Branch Entry', html)

    def test_full_report_is_unchanged_for_superusers_and_self(self):
        admin = User.objects.create_superuser('root', 'r@example.com', 'rootpass123')
        self.assertEqual(user_report_access(admin, self.worker), USER_REPORT_FULL)
        self.assertEqual(user_report_access(self.worker, self.worker), USER_REPORT_FULL)
        client = Client()
        client.login(username='root', password='rootpass123')
        html = json.loads(client.get(
            reverse('user_report_modal', args=[self.worker.pk]), {'window': 'all'},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        ).content)['html']
        self.assertIn('worker@example.com', html)

    def test_staff_without_report_permission_is_still_refused(self):
        plain = User.objects.create_user('plainstaff', password='plainpass1', is_staff=True)
        self.assertIsNone(user_report_access(plain, self.worker))
        client = Client()
        client.login(username='plainstaff', password='plainpass1')
        self.assertEqual(client.get(reverse('user_report_modal', args=[self.worker.pk])).status_code, 403)

    def test_every_user_the_reports_page_lists_opens(self):
        page = self.client.get(reverse('reports_overview'), {'window': 'all'})
        self.assertEqual(page.status_code, 200)
        listed = page.context['overview']['visible_users']
        self.assertTrue(listed)
        for user in listed:
            self.assertEqual(self._modal(user).status_code, 200, user.username)
